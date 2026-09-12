# Copyright (c) 2026 ICTU-CRIS contributors
# SPDX-License-Identifier: Apache-2.0
"""Bản báo cáo kỳ đóng băng, có phiên bản (`cris.report`, lát cắt L1): nghiệp
vụ trực tiếp trên PostgreSQL thật (fixture `conn`, theo khuôn
`test_declare.py`) và lớp API mỏng `cris/api/routes/reports.py` (theo khuôn
`test_api_declarations.py`)."""
import io
import uuid
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
from fastapi.testclient import TestClient

from cris import declare, period, report, rules
from cris.ai.provider import clear_provider_cache
from cris.api.app import create_app


def q(conn, sql, *a):
    with conn.cursor() as cur:
        cur.execute(sql, a)
        return cur.fetchall() if cur.description else None


def mk_unit(conn, code="khoa-cntt", name="Khoa Công nghệ thông tin"):
    return q(conn, "INSERT INTO unit(code, name) VALUES (%s,%s) RETURNING id", code, name)[0]["id"]


def mk_actor(conn, roles, unit_id=None, email=None):
    """Tạo `app_user` và commit ngay — dùng qua header `X-CRIS-User` bởi
    `TestClient`, tức một kết nối khác của `conn`, nên phải thấy được hàng
    vừa tạo dù không có lời gọi nào khác commit hộ sau đó."""
    email = email or f"actor-{uuid.uuid4().hex[:10]}@ictu.test"
    uid = q(conn, "INSERT INTO app_user(email, display_name, roles, unit_id) VALUES (%s,%s,%s,%s) RETURNING id",
            email, "Người kiểm thử L1", roles, unit_id)[0]["id"]
    conn.commit()
    return uid


def mk_work(conn, title="Bài báo báo cáo", doc_type="bai_bao", doi=None, year=2026,
            quartile=None, journal=None, indexes=None):
    q(conn, "INSERT INTO sync_run(source, scope) VALUES ('manual','t')")
    q(conn, "INSERT INTO source_record(sync_run_id, source, source_key, doc_type, content_hash, raw) "
            "VALUES (currval('sync_run_id_seq'),'manual',%s,%s,'h','{}')", title, doc_type)
    wid = q(conn,
            "INSERT INTO work(doc_type, primary_source_record_id, title, title_norm, doi, year_issue, "
            "quartile, journal, indexes, state) "
            "VALUES (%s, currval('source_record_id_seq'), %s, %s, %s, %s, %s, %s, %s, 'DaChuanHoa') RETURNING id",
            doc_type, title, title.lower(), doi, year, quartile, journal, indexes or [])[0]["id"]
    conn.commit()
    return wid


def mk_mention(conn, work_id, raw_name, role="author", position=1):
    return q(conn,
             "INSERT INTO author_mention(work_id, role, position, raw_name, name_norm, name_key) "
             "VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",
             work_id, role, position, raw_name, raw_name.lower(), raw_name.lower())[0]["id"]


def mk_person(conn, name, unit_id=None):
    return q(conn,
             "INSERT INTO person(kind, display_name, name_norm, name_keys, unit_id) "
             "VALUES ('lecturer',%s,%s,%s,%s) RETURNING id",
             name, name.lower(), [name.lower()], unit_id)[0]["id"]


def mk_link(conn, mention_id, person_id, state="DaXacNhan"):
    q(conn, "INSERT INTO author_link(mention_id, person_id, confidence, state) "
            "VALUES (%s,%s,'ten_day_du_duy_nhat',%s)", mention_id, person_id, state)
    conn.commit()


def due_soon(days=30):
    return datetime.now(UTC) + timedelta(days=days)


def open_basic(conn, actor_id, code="K2026-REP"):
    return period.open_period(conn, code=code, name="Kỳ báo cáo", scope={"doc_types": ["bai_bao"]},
                              criteria=None, due_at=due_soon(), actor_id=actor_id)


def as_user(uid):
    return {"X-CRIS-User": str(uid)}


@pytest.fixture(autouse=True)
def seed(conn):
    rules.seed_rules(conn, None)


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("CRIS_AI_PROVIDER", "fake")
    clear_provider_cache()
    return TestClient(create_app(static_dir="/nonexistent"))


# ---------- nghiệp vụ: cris.report ----------

def test_build_report_version_increments(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-1")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    r1 = report.build_report(conn, pid, user_id)
    r2 = report.build_report(conn, pid, user_id)
    assert report.get_report(conn, r1)["version"] == 1
    assert report.get_report(conn, r2)["version"] == 2


def test_build_report_sha_stable_without_data_change(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-2")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    r1 = report.build_report(conn, pid, user_id)
    r2 = report.build_report(conn, pid, user_id)
    assert report.get_report(conn, r1)["sha256"] == report.get_report(conn, r2)["sha256"]


def test_old_report_payload_frozen_after_state_change(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-3")
    did = declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    r1 = report.build_report(conn, pid, user_id)
    before = report.get_report(conn, r1)
    payload_before, sha_before = before["payload"], before["sha256"]

    declare.set_state(conn, did, "ChoKhoaDuyet", user_id)

    after = report.get_report(conn, r1)
    assert after["payload"] == payload_before
    assert after["sha256"] == sha_before
    assert after["payload"]["items"][0]["state"] == "Nhap"

    r2 = report.build_report(conn, pid, user_id)
    fresh = report.get_report(conn, r2)
    assert fresh["sha256"] != sha_before
    assert fresh["payload"]["items"][0]["state"] == "ChoKhoaDuyet"


def test_build_report_empty_period_has_no_items(conn, user_id):
    pid = open_basic(conn, user_id, "K2026-REP-4")
    rid = report.build_report(conn, pid, user_id)
    row = report.get_report(conn, rid)
    assert row["payload"]["items"] == []
    assert row["summary"]["totals"] == {"declared": 0, "accepted": 0, "by_state": {}}
    assert row["summary"]["units"] == []


def test_build_report_authors_mix_linked_and_raw(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn, title="Bài có tác giả")
    m1 = mk_mention(conn, work_id, "Nguyen Van A", position=1)
    mk_mention(conn, work_id, "Tran Thi B", position=2)
    person_id = mk_person(conn, "PGS.TS Nguyễn Văn A")
    mk_link(conn, m1, person_id, state="DaXacNhan")
    pid = open_basic(conn, user_id, "K2026-REP-5")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    item = report.get_report(conn, rid)["payload"]["items"][0]
    assert item["authors"] == ["PGS.TS Nguyễn Văn A", "Tran Thi B"]


def test_build_report_evidence_count(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-6")
    did = declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    declare.add_evidence(conn, did, kind="note", note="ghi chú", actor_id=user_id)
    declare.add_evidence(conn, did, kind="link", url="http://x.test", actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    item = report.get_report(conn, rid)["payload"]["items"][0]
    assert item["evidence_count"] == 2


def test_build_report_events_are_reduced(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-7")
    did = declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    declare.set_state(conn, did, "ChoKhoaDuyet", user_id)

    rid = report.build_report(conn, pid, user_id)
    item = report.get_report(conn, rid)["payload"]["items"][0]
    assert [e["state"] for e in item["events"]] == ["Nhap", "ChoKhoaDuyet"]
    assert set(item["events"][0]) == {"state", "at"}


def test_build_report_summary_by_unit_and_doc_type(conn, user_id):
    u1 = mk_unit(conn, "khoa-a", "Khoa A")
    u2 = mk_unit(conn, "khoa-b", "Khoa B")
    w1 = mk_work(conn, "Bài 1", doc_type="bai_bao")
    w2 = mk_work(conn, "Đồ án 1", doc_type="do_an")
    pid = open_basic(conn, user_id, "K2026-REP-8")
    declare.add_declaration(conn, period_id=pid, work_id=w1, unit_id=u1, actor_id=user_id)
    declare.add_declaration(conn, period_id=pid, work_id=w2, unit_id=u2, actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    summary = report.get_report(conn, rid)["summary"]
    declared_by_code = {u["code"]: u["declared"] for u in summary["units"]}
    assert declared_by_code == {"khoa-a": 1, "khoa-b": 1}
    assert summary["by_doc_type"] == {"bai_bao": 1, "do_an": 1}
    assert summary["totals"]["declared"] == 2
    assert summary["totals"]["accepted"] == 0


def test_finalize_period_creates_report(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-9")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    period.close_submissions(conn, pid, user_id)

    result = declare.finalize_period(conn, pid, user_id)
    assert "report_id" in result
    row = report.get_report(conn, result["report_id"])
    assert row is not None
    assert row["period_id"] == pid
    assert row["version"] == 1
    assert row["note"] == "Tự sinh khi chốt kỳ"


def test_report_audit_log_entry(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-10")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    entry = q(conn, "SELECT action, entity, entity_id FROM audit_log WHERE action='period.report' "
                    "AND entity_id=%s", rid)
    assert entry and entry[0]["entity"] == "period_report"


# ---------- xuất CSV/XLSX ----------

def test_report_csv_has_bom_and_row_per_item(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn, title="Bài CSV", year=2025, doi="10.1/x", quartile="Q1",
                      journal="Tạp chí A", indexes=["scopus"])
    pid = open_basic(conn, user_id, "K2026-REP-11")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    row = report.get_report(conn, rid)
    text = report.report_csv(row)
    assert text.startswith("﻿")
    lines = text[1:].strip("\r\n").split("\r\n")
    assert len(lines) == 2
    assert "Bài CSV" in lines[1] and "Q1" in lines[1] and "scopus" in lines[1]


def test_report_xlsx_has_two_sheets_and_row_count(conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn, title="Bài XLSX")
    pid = open_basic(conn, user_id, "K2026-REP-12")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    rid = report.build_report(conn, pid, user_id)
    row = report.get_report(conn, rid)
    data = report.report_xlsx(row)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Tổng hợp", "Chi tiết"]
    ws_detail = wb["Chi tiết"]
    assert ws_detail.max_row == 2  # tiêu đề + 1 hồ sơ
    ws_summary = wb["Tổng hợp"]
    assert ws_summary["A1"].value == "Đơn vị"
    assert ws_summary.freeze_panes == "A2"
    assert ws_detail.freeze_panes == "A2"


# ---------- API: cris/api/routes/reports.py ----------

def test_api_create_and_list_reports(client, conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-13")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)

    r = client.post(f"/api/periods/{pid}/reports", json={"note": "kiểm thử"})
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["version"] == 1 and body["note"] == "kiểm thử"
    assert body["summary"]["totals"]["declared"] == 1

    listing = client.get(f"/api/periods/{pid}/reports").json()
    assert len(listing) == 1 and listing[0]["id"] == body["id"]


def test_api_create_report_unknown_period_is_404(client, user_id):
    r = client.post("/api/periods/999999/reports", json={})
    assert r.status_code == 404


def test_api_create_report_requires_rd_officer(client, conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-14")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    faculty_actor = mk_actor(conn, ["faculty_officer"], unit_id=unit_id)

    r = client.post(f"/api/periods/{pid}/reports", json={}, headers=as_user(faculty_actor))
    assert r.status_code == 403


def test_api_report_detail_scoped_by_unit_for_faculty(client, conn, user_id):
    u1 = mk_unit(conn, "khoa-a", "Khoa A")
    u2 = mk_unit(conn, "khoa-b", "Khoa B")
    w1 = mk_work(conn, "Bài A")
    w2 = mk_work(conn, "Bài B")
    pid = open_basic(conn, user_id, "K2026-REP-15")
    declare.add_declaration(conn, period_id=pid, work_id=w1, unit_id=u1, actor_id=user_id)
    declare.add_declaration(conn, period_id=pid, work_id=w2, unit_id=u2, actor_id=user_id)
    rid = client.post(f"/api/periods/{pid}/reports", json={}).json()["id"]

    faculty_head = mk_actor(conn, ["faculty_head"], unit_id=u1)
    detail = client.get(f"/api/reports/{rid}", headers=as_user(faculty_head)).json()
    assert len(detail["items"]) == 1
    assert detail["items"][0]["unit"]["code"] == "khoa-a"
    assert detail["summary"]["totals"]["declared"] == 2  # tổng không bị lọc theo đơn vị

    rd_view = client.get(f"/api/reports/{rid}", headers=as_user(user_id)).json()
    assert len(rd_view["items"]) == 2


def test_api_report_detail_404(client, user_id):
    assert client.get("/api/reports/999999").status_code == 404


def test_api_report_export_unknown_format_is_400(client, conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-16")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    rid = client.post(f"/api/periods/{pid}/reports", json={}).json()["id"]

    r = client.get(f"/api/reports/{rid}/export", params={"format": "pdf"})
    assert r.status_code == 400


def test_api_report_export_csv_and_xlsx(client, conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-17")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    created = client.post(f"/api/periods/{pid}/reports", json={}).json()
    rid = created["id"]
    period_code = "K2026-REP-17"

    csv_r = client.get(f"/api/reports/{rid}/export", params={"format": "csv"})
    assert csv_r.status_code == 200
    assert csv_r.headers["content-disposition"] == \
        f'attachment; filename="bao-cao-{period_code}-v{created["version"]}.csv"'
    assert csv_r.content.startswith("﻿".encode("utf-8"))

    xlsx_r = client.get(f"/api/reports/{rid}/export", params={"format": "xlsx"})
    assert xlsx_r.status_code == 200
    assert xlsx_r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert xlsx_r.headers["content-disposition"] == \
        f'attachment; filename="bao-cao-{period_code}-v{created["version"]}.xlsx"'


def test_api_finalize_returns_report_id(client, conn, user_id):
    unit_id = mk_unit(conn)
    work_id = mk_work(conn)
    pid = open_basic(conn, user_id, "K2026-REP-18")
    declare.add_declaration(conn, period_id=pid, work_id=work_id, unit_id=unit_id, actor_id=user_id)
    assert client.post(f"/api/periods/{pid}/close").status_code == 200

    r = client.post(f"/api/periods/{pid}/finalize")
    assert r.status_code == 200, r.text
    body = r.json()
    assert "report_id" in body and body["report_id"]

    detail = client.get(f"/api/reports/{body['report_id']}").json()
    assert detail["note"] == "Tự sinh khi chốt kỳ"
