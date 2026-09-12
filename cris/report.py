# Copyright (c) 2026 ICTU-CRIS contributors
# SPDX-License-Identifier: Apache-2.0
"""Bản báo cáo kỳ đóng băng, có phiên bản (lát cắt L1 — điểm BA "P": phiên bản
báo cáo được đóng băng sau khi chốt). `build_report` chụp lại TOÀN BỘ hồ sơ kê
khai của một kỳ (mọi trạng thái) cùng công trình/tác giả/minh chứng tại thời
điểm gọi, ghi thành một dòng `period_report` mới (migration
`0018_period_report.sql`) — payload đã ghi không bao giờ bị sửa lại, kể cả khi
declaration/work/author_link đổi tiếp sau đó (khác các view/API tra cứu khác
luôn đọc trạng thái sống). `version` tăng dần theo từng kỳ; `sha256` băm JSON
canonical (`sort_keys=True, ensure_ascii=False`) của `payload` để bất kỳ ai
cũng kiểm lại được là báo cáo không bị sửa sau khi sinh.

Tác giả của mỗi công trình: tái dùng thứ tự/lọc vai trò của `cris.cite.authors_of`
(sinh viên/tác giả trước, GVHD sau, theo `position`), rồi with mỗi lượt tên
tra thêm liên kết còn sống (`author_link.state IN ('DaNoiTuDong','DaXacNhan')`)
— có thì dùng `person.display_name`, không thì giữ `raw_name` (tên thô).

Xuất báo cáo: `report_csv` (stdlib `csv`, BOM UTF-8) và `report_xlsx`
(`openpyxl`, 2 sheet) đọc từ một report đã lấy bằng `get_report` (hoặc một bản
đã lọc `payload["items"]` theo đơn vị — xem `cris/api/routes/reports.py`), nên
không tự truy vấn DB.
"""
import hashlib
import io
import json

from cris import audit, cite
from cris.db import tx
from cris.period import DECLARATION_STATES

DOC_TYPE_LABELS = {
    "bai_bao": "Bài báo", "do_an": "Đồ án", "luan_van": "Luận văn", "luan_an": "Luận án",
    "hoc_lieu": "Học liệu",
}

STATE_LABELS = {
    "Nhap": "Nháp",
    "ChoBoSung": "Chờ bổ sung",
    "ChoKhoaDuyet": "Chờ khoa duyệt",
    "KhoaDaDuyet": "Khoa đã duyệt",
    "ChoPhongKiemTra": "Chờ phòng kiểm tra",
    "DatYeuCau": "Đạt yêu cầu",
    "DaChot": "Đã chốt",
    "Rut": "Đã rút",
}

CSV_HEADER = ["Đơn vị", "Trạng thái", "Loại", "Tiêu đề", "Tác giả", "Năm", "DOI", "Chỉ mục", "Quartile", "Số minh chứng"]


def _authors_for_work(conn, work_id):
    """Danh sách tên tác giả của một công trình theo đúng thứ tự/lọc vai trò
    của `cite.authors_of`: mỗi lượt tên đã có liên kết còn sống thì thay bằng
    tên hiển thị của người đó, còn lại giữ nguyên tên thô."""
    ordered = cite.authors_of(conn, work_id)
    if not ordered:
        return []
    with conn.cursor() as cur:
        cur.execute(
            "SELECT m.role, m.position, p.display_name FROM author_mention m "
            "JOIN author_link l ON l.mention_id = m.id AND l.state IN ('DaNoiTuDong','DaXacNhan') "
            "JOIN person p ON p.id = l.person_id "
            "WHERE m.work_id=%s AND m.position > 0 AND m.role IN ('author','student','mentor')",
            (work_id,))
        linked = {(r["role"], r["position"]): r["display_name"] for r in cur.fetchall()}
    return [linked.get((r["role"], r["position"]), r["raw_name"]) for r in ordered]


def _row_to_item(conn, d):
    with conn.cursor() as cur:
        cur.execute("SELECT count(*) AS n FROM evidence WHERE declaration_id=%s", (d["declaration_id"],))
        evidence_count = cur.fetchone()["n"]
        cur.execute("SELECT to_state, at FROM declaration_event WHERE declaration_id=%s ORDER BY at, id",
                    (d["declaration_id"],))
        events = [{"state": e["to_state"], "at": e["at"].isoformat()} for e in cur.fetchall()]
    return {
        "declaration_id": d["declaration_id"],
        "state": d["state"],
        "unit": {"id": d["unit_id"], "code": d["unit_code"], "name": d["unit_name"]},
        "work": {
            "id": d["work_id"], "title": d["title"], "doc_type": d["doc_type"],
            "year_issue": d["year_issue"], "doi": d["doi"],
            "indexes": list(d["indexes"] or []), "quartile": d["quartile"], "journal": d["journal"],
        },
        "authors": _authors_for_work(conn, d["work_id"]),
        "evidence_count": evidence_count,
        "events": events,
    }


def build_report(conn, period_id, actor_id, note=None):
    """Sinh một bản báo cáo mới (phiên bản kế tiếp) cho một kỳ báo cáo, chụp
    lại TOÀN BỘ hồ sơ kê khai của kỳ (mọi trạng thái) — kể cả `Nhap`/`Rut`,
    không chỉ hồ sơ đã chốt, để báo cáo phản ánh đúng bức tranh kê khai tại
    thời điểm sinh. `ValueError` nếu không tìm thấy kỳ. Ghi
    `audit_log('period.report')`. Trả về id của bản báo cáo vừa tạo."""
    with tx(conn), conn.cursor() as cur:
        cur.execute("SELECT id FROM period WHERE id=%s FOR UPDATE", (period_id,))
        if cur.fetchone() is None:
            raise ValueError(f"không tìm thấy kỳ báo cáo: {period_id}")

        cur.execute(
            "SELECT d.id AS declaration_id, d.state, d.unit_id, u.code AS unit_code, u.name AS unit_name, "
            "d.work_id, w.title, w.doc_type, w.year_issue, w.doi, w.indexes, w.quartile, w.journal "
            "FROM declaration d JOIN unit u ON u.id = d.unit_id JOIN work w ON w.id = d.work_id "
            "WHERE d.period_id=%s ORDER BY d.id", (period_id,))
        decl_rows = cur.fetchall()

        items = [_row_to_item(conn, d) for d in decl_rows]

        by_unit = {}
        by_doc_type = {}
        total_by_state = {}
        for it in items:
            u = by_unit.setdefault(it["unit"]["id"], {
                "unit_id": it["unit"]["id"], "code": it["unit"]["code"], "name": it["unit"]["name"],
                "by_state": {}, "declared": 0, "accepted": 0})
            u["by_state"][it["state"]] = u["by_state"].get(it["state"], 0) + 1
            u["declared"] += 1
            if it["state"] == "DaChot":
                u["accepted"] += 1
            by_doc_type[it["work"]["doc_type"]] = by_doc_type.get(it["work"]["doc_type"], 0) + 1
            total_by_state[it["state"]] = total_by_state.get(it["state"], 0) + 1

        summary = {
            "units": sorted(by_unit.values(), key=lambda u: u["code"]),
            "by_doc_type": by_doc_type,
            "totals": {
                "declared": len(items),
                "accepted": total_by_state.get("DaChot", 0),
                "by_state": total_by_state,
            },
        }
        payload = {"items": items}
        sha256 = hashlib.sha256(
            json.dumps(payload, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()

        cur.execute("SELECT COALESCE(max(version), 0) AS v FROM period_report WHERE period_id=%s", (period_id,))
        version = cur.fetchone()["v"] + 1

        cur.execute(
            "INSERT INTO period_report(period_id, version, generated_by, note, summary, payload, sha256) "
            "VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id",
            (period_id, version, actor_id, note,
             json.dumps(summary, ensure_ascii=False), json.dumps(payload, ensure_ascii=False), sha256))
        report_id = cur.fetchone()["id"]
        audit.log(conn, actor_id, "period.report", "period_report", report_id,
                  after={"period_id": period_id, "version": version, "sha256": sha256})
    return report_id


def get_report(conn, report_id):
    """Một bản báo cáo đầy đủ (kèm `period_code`/`period_name` để đặt tên tệp
    xuất) theo id, hoặc `None` nếu không có."""
    with conn.cursor() as cur:
        cur.execute(
            "SELECT r.*, p.code AS period_code, p.name AS period_name "
            "FROM period_report r JOIN period p ON p.id = r.period_id WHERE r.id=%s", (report_id,))
        return cur.fetchone()


def list_reports(conn, period_id):
    """Các phiên bản báo cáo của một kỳ, mới nhất trước; mỗi dòng kèm `totals`
    lấy từ `summary` (không tải `payload` đầy đủ)."""
    with conn.cursor() as cur:
        cur.execute(
            "SELECT r.id, r.version, r.generated_at, r.note, r.summary, u.display_name AS generated_by_name "
            "FROM period_report r LEFT JOIN app_user u ON u.id = r.generated_by "
            "WHERE r.period_id=%s ORDER BY r.version DESC", (period_id,))
        rows = cur.fetchall()
    return [{
        "id": r["id"], "version": r["version"], "generated_at": r["generated_at"],
        "generated_by_name": r["generated_by_name"], "note": r["note"], "totals": r["summary"]["totals"],
    } for r in rows]


def report_csv(report):
    """CSV UTF-8 có BOM (Excel mở đúng dấu tiếng Việt) từ `items` của một
    report (đã lấy bằng `get_report`, có thể đã lọc theo đơn vị)."""
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_HEADER)
    for it in report["payload"]["items"]:
        w = it["work"]
        writer.writerow([
            it["unit"]["name"],
            STATE_LABELS.get(it["state"], it["state"]),
            DOC_TYPE_LABELS.get(w["doc_type"], w["doc_type"]),
            w["title"] or "",
            "; ".join(it["authors"]),
            w["year_issue"] if w["year_issue"] is not None else "",
            w["doi"] or "",
            ", ".join(w["indexes"] or []),
            w["quartile"] or "",
            it["evidence_count"],
        ])
    return "﻿" + buf.getvalue()


def _autosize_columns(ws, max_width=60):
    from openpyxl.utils import get_column_letter
    for col_cells in ws.columns:
        col_cells = list(col_cells)
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, 8), max_width)


def report_xlsx(report):
    """XLSX (`openpyxl`) từ một report (đã lấy bằng `get_report`, có thể đã
    lọc theo đơn vị): sheet "Tổng hợp" (đơn vị × trạng thái, rồi tổng theo
    loại tài liệu) và sheet "Chi tiết" (từng hồ sơ). Hàng đầu mỗi sheet in
    đậm, cố định (`freeze_panes="A2"`), cột rộng theo nội dung (tối đa 60)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    summary = report["summary"]
    items = report["payload"]["items"]
    bold = Font(bold=True)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Tổng hợp"
    ws1.append(["Đơn vị", *[STATE_LABELS.get(s, s) for s in DECLARATION_STATES], "Tổng"])
    for cell in ws1[1]:
        cell.font = bold
    for u in summary["units"]:
        ws1.append([u["name"], *[u["by_state"].get(s, 0) for s in DECLARATION_STATES], u["declared"]])
    ws1.append([])
    header_row = ws1.max_row + 1
    ws1.append(["Loại", "Số lượng"])
    for cell in ws1[header_row]:
        cell.font = bold
    for doc_type, n in sorted(summary["by_doc_type"].items()):
        ws1.append([DOC_TYPE_LABELS.get(doc_type, doc_type), n])
    ws1.freeze_panes = "A2"
    _autosize_columns(ws1)

    ws2 = wb.create_sheet("Chi tiết")
    ws2.append(CSV_HEADER)
    for cell in ws2[1]:
        cell.font = bold
    for it in items:
        w = it["work"]
        ws2.append([
            it["unit"]["name"],
            STATE_LABELS.get(it["state"], it["state"]),
            DOC_TYPE_LABELS.get(w["doc_type"], w["doc_type"]),
            w["title"] or "",
            "; ".join(it["authors"]),
            w["year_issue"] if w["year_issue"] is not None else "",
            w["doi"] or "",
            ", ".join(w["indexes"] or []),
            w["quartile"] or "",
            it["evidence_count"],
        ])
    ws2.freeze_panes = "A2"
    _autosize_columns(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
